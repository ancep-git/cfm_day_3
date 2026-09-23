"""Builds Year1_Ball_Bearing_PL.xlsx: contracts won in the Year 1 auction -> contract profitability -> P&L."""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F = "Arial"
BLUE, GREEN = Font(name=F, color="0000FF"), Font(name=F, color="008000")
BOLD, NORM = Font(name=F, bold=True), Font(name=F)
TITLE = Font(name=F, bold=True, size=14)
YELLOW = PatternFill("solid", start_color="FFFF00")
HDR = PatternFill("solid", start_color="D9E1F2")
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
INT = '#,##0;(#,##0);"-"'
TOP = Border(top=Side(style="thin"))

wb = Workbook()

# ---------------- Inputs ----------------
s = wb.active
s.title = "Inputs"
s["A1"] = "Ball Bearing Contract Auction - Inputs"; s["A1"].font = TITLE
s["A2"] = "Blue = hard-coded input from the case brief. Yellow = decision/assumption you can change."
s["A2"].font = Font(name=F, italic=True)

s["A4"] = "Raw-material price per BB (applies to every BB in one purchase order)"; s["A4"].font = BOLD
for c, h in zip("ABCDE", ["Qty from", "Qty to", "Blue", "Silver", "Copper"]):
    s[f"{c}5"] = h; s[f"{c}5"].font = BOLD; s[f"{c}5"].fill = HDR
tiers = [(1, 10, 7000, 6500, 5500), (11, 20, 6500, 6000, 5000), (21, 30, 5500, 5500, 4000),
         (31, 40, 5000, 4500, 3000), (41, "or more", 4000, 3500, 2000)]
for i, row in enumerate(tiers, start=6):
    for c, v in zip("ABCDE", row):
        s[f"{c}{i}"] = v; s[f"{c}{i}"].font = BLUE
        if c in "CDE": s[f"{c}{i}"].number_format = USD
s["A11"] = "Source: case brief, raw-material price table. Gold ($4,500-$5,500) is announced after the auction; no Gold contract was won."
s["A11"].font = Font(name=F, italic=True, size=9)

params = [
    (13, "Rejection rate", 0.15, PCT, "Case brief: 15% of production is rejected"),
    (14, "Wage coefficient (x raw materials)", 1.8, "0.0x", "Case brief: wages = raw materials x 1.8"),
    (15, "Other direct contract expenses per contract", 0, USD, "None given in the case brief"),
    (16, "Combine same-spec contracts into one purchase order? (1 = yes, 0 = no)", 1, "0",
     "Assumption: the brief says tier price applies to every BB in a purchase order and separate orders are "
     "priced separately, so one Silver order for both Silver contracts is allowed. Set 0 to price each contract separately."),
    (17, "Year 1 market size", 5300000, USD, "Case brief: market outlook"),
]
for r, label, v, fmt, note in params:
    s[f"A{r}"] = label; s[f"A{r}"].font = NORM
    s[f"D{r}"] = v; s[f"D{r}"].font = BLUE; s[f"D{r}"].number_format = fmt
    s[f"D{r}"].comment = Comment(note, "Case")
s["D16"].fill = YELLOW

s["A19"] = "Annual fixed operating expenses"; s["A19"].font = BOLD
fixed = [("Factory rent", 50000), ("Management and administration", 40000), ("Production-machine maintenance", 35000),
         ("Utilities and insurance", 20000), ("Office and marketing", 15000), ("Depreciation", 40000)]
for i, (k, v) in enumerate(fixed, start=20):
    s[f"A{i}"] = k; s[f"A{i}"].font = NORM
    s[f"D{i}"] = v; s[f"D{i}"].font = BLUE; s[f"D{i}"].number_format = USD
s["A26"] = "Total fixed operating expenses"; s["A26"].font = BOLD
s["D26"] = "=SUM(D20:D25)"; s["D26"].font = BOLD; s["D26"].number_format = USD; s["D26"].border = TOP
s.column_dimensions["A"].width = 16
for c in "BCDE": s.column_dimensions[c].width = 13

# ---------------- Contracts ----------------
c = wb.create_sheet("Contracts")
c["A1"] = "Year 1 contracts won (from the auction cards) and contract profitability"; c["A1"].font = TITLE
c["A2"] = "Blue = values read from the won contract cards (handwritten winning price). All other cells are formulas."
c["A2"].font = Font(name=F, italic=True)
heads = ["Contract", "Spec", "Qty accepted (BB)", "Winning price / BB", "Revenue", "BB to produce",
         "Purchase-order qty", "Raw-mat. price / BB", "Raw materials", "Production wages",
         "Other direct exp.", "Contribution margin", "CM %", "Break-even price / BB",
         "List price / BB", "Discount vs list"]
for i, h in enumerate(heads):
    cell = c.cell(row=4, column=i + 1, value=h)
    cell.font = BOLD; cell.fill = HDR; cell.alignment = Alignment(wrap_text=True, vertical="center")
c.row_dimensions[4].height = 45
won = [("05", "Silver", 15, 18000, 20000), ("08", "Silver", 40, 11000, 20000), ("09", "Copper", 40, 6500, 18000)]
first, last = 5, 4 + len(won)
for r, (no, spec, q, p, lp) in enumerate(won, start=first):
    for col, v in zip("ABCDO", [no, spec, q, p, lp]):
        c[f"{col}{r}"] = v; c[f"{col}{r}"].font = BLUE
    c[f"E{r}"] = f"=C{r}*D{r}"
    c[f"F{r}"] = f"=ROUNDUP(ROUND(C{r}/(1-Inputs!$D$13),6),0)"
    c[f"G{r}"] = f"=IF(Inputs!$D$16=1,SUMIF($B${first}:$B${last},B{r},$F${first}:$F${last}),F{r})"
    c[f"H{r}"] = f"=INDEX(Inputs!$C$6:$E$10,MATCH(G{r},Inputs!$A$6:$A$10,1),MATCH(B{r},Inputs!$C$5:$E$5,0))"
    c[f"I{r}"] = f"=F{r}*H{r}"
    c[f"J{r}"] = f"=I{r}*Inputs!$D$14"
    c[f"K{r}"] = "=Inputs!$D$15"
    c[f"L{r}"] = f"=E{r}-I{r}-J{r}-K{r}"
    c[f"M{r}"] = f"=IF(E{r}=0,0,L{r}/E{r})"
    c[f"N{r}"] = f"=(I{r}+J{r}+K{r})/C{r}"
    c[f"P{r}"] = f"=IF(O{r}=0,0,1-D{r}/O{r})"
    for col in "HK": c[f"{col}{r}"].font = GREEN
tr = last + 1
c[f"A{tr}"] = "Total"; c[f"A{tr}"].font = BOLD
for col in "CEFIJKL":
    c[f"{col}{tr}"] = f"=SUM({col}{first}:{col}{last})"
c[f"M{tr}"] = f"=IF(E{tr}=0,0,L{tr}/E{tr})"
c[f"D{tr}"] = f"=IF(C{tr}=0,0,E{tr}/C{tr})"
c[f"N{tr}"] = f"=(I{tr}+J{tr}+K{tr})/C{tr}"
for col in "ABCDEFGHIJKLMNOP":
    c[f"{col}{tr}"].border = TOP
    if col != "A" and c[f"{col}{tr}"].value is not None: c[f"{col}{tr}"].font = BOLD
for r in range(first, tr + 1):
    for col in "DEHIJKLNO": c[f"{col}{r}"].number_format = USD
    for col in "CFG": c[f"{col}{r}"].number_format = INT
    for col in "MP": c[f"{col}{r}"].number_format = PCT
notes = [
    "Notes",
    "BB to produce = accepted qty / (1 - 15%), rounded up to the next whole BB (rejects still cost materials and wages).",
    "Purchase-order qty decides the raw-material price tier. With Inputs!D16 = 1 both Silver contracts share one order (18 + 48 = 66 BB).",
    "Break-even price = the lowest winning price per accepted BB that covers raw materials + wages (before fixed costs).",
    "Contract 09 card: the quantity line is struck through along with the price; quantity is taken as the printed 40 BB.",
]
for i, t in enumerate(notes, start=tr + 2):
    c[f"A{i}"] = t; c[f"A{i}"].font = BOLD if i == tr + 2 else Font(name=F, size=9)
c.column_dimensions["A"].width = 10
for col in "BCDEFGHIJKLMNOP": c.column_dimensions[col].width = 13
c.freeze_panes = "C5"

# ---------------- P&L ----------------
p = wb.create_sheet("P&L")
p["A1"] = "Annual profit and loss statement"; p["A1"].font = TITLE
p["B3"] = "Year 1"; p["B3"].font = BOLD; p["B3"].fill = HDR; p["A3"].fill = HDR
rows = [
    (4, "Contract revenue", f"=Contracts!E{tr}"),
    (5, "Less raw materials used", f"=-Contracts!I{tr}"),
    (6, "Less production wages", f"=-Contracts!J{tr}"),
    (7, "Less other direct contract expenses", f"=-Contracts!K{tr}"),
    (8, "Contribution margin", "=SUM(B4:B7)"),
    (9, "Less fixed operating expenses", "=-Inputs!D26"),
    (10, "Operating profit", "=B8+B9"),
    (12, "Contribution-margin %", "=IF(B4=0,0,B8/B4)"),
    (13, "Operating-margin %", "=IF(B4=0,0,B10/B4)"),
    (14, "Cumulative operating profit", "=B10"),
    (16, "BB accepted by customers", f"=Contracts!C{tr}"),
    (17, "BB produced (incl. rejects)", f"=Contracts!F{tr}"),
    (18, "Share of Year 1 market (revenue)", "=IF(Inputs!D17=0,0,B4/Inputs!D17)"),
    (19, "Operating profit per accepted BB", "=IF(B16=0,0,B10/B16)"),
]
for r, label, f in rows:
    p[f"A{r}"] = label; p[f"B{r}"] = f
    bold = r in (8, 10, 14)
    p[f"A{r}"].font = BOLD if bold else NORM
    p[f"B{r}"].font = GREEN if r in (4, 5, 6, 7, 9, 16, 17) else (BOLD if bold else NORM)
    p[f"B{r}"].number_format = PCT if r in (12, 13, 18) else (INT if r in (16, 17) else USD)
    if r in (8, 10): p[f"B{r}"].border = TOP
p["A21"] = "Green = linked from another sheet. Change Inputs!D16 to 0 to see the P&L if every contract is ordered separately."
p["A21"].font = Font(name=F, italic=True, size=9)
p.column_dimensions["A"].width = 40; p.column_dimensions["B"].width = 16

wb.calculation.fullCalcOnLoad = True
wb.save("/home/user/cfm_day_3/ball-bearing-auction/Year1_Ball_Bearing_PL.xlsx")
