"""Recalculate BB_Budget.xlsx with LibreOffice (headless) and check the expected results.

Works on temporary copies, so the delivered workbook is never modified.
Usage:  python check_workbook.py [path/to/BB_Budget.xlsx]
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("BB_Budget.xlsx")
ERRORS = ("#REF!", "#N/A", "#VALUE!", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!")
failures = []


def recalc(path):
    out = path.parent / "out"
    out.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(out), str(path)], check=True, capture_output=True)
    return load_workbook(out / path.name, data_only=True)


def scan_errors(wb, tag):
    found = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets
             for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value in ERRORS]
    print(f"[{tag}] formula errors: {len(found)}" + (f" -> {found[:20]}" if found else ""))
    failures.extend(found)


def check(label, got, want, tol=0.0005):
    ok = got is not None and abs(got - want) <= tol
    print(f"  {'OK  ' if ok else 'FAIL'} {label:<38} got {got!r:>14}  want {want!r}")
    if not ok:
        failures.append(label)


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)

    # --- Base case --------------------------------------------------------
    base = tmp / "base.xlsx"
    shutil.copy(SRC, base)
    wb = recalc(base)
    scan_errors(wb, "base")
    cs, pl, sm = wb["Contracts Y1"], wb["P&L Y1"], wb["Summary"]
    rows = {cs[f"A{r}"].value: r for r in range(5, 35) if cs[f"A{r}"].value}
    want = {
        "C05": dict(E=18, F=3, G=6000, H=270000, I=108000, J=194400, L=-32400, M=-0.12, N=20160),
        "C08": dict(E=48, F=8, G=3500, H=440000, I=168000, J=302400, L=-30400, M=-0.0691, N=11760),
        "C09": dict(E=48, F=8, G=2000, H=260000, I=96000, J=172800, L=-8800, M=-0.0338, N=6720),
    }
    for cid, cols in want.items():
        print(cid)
        for col, v in cols.items():
            check(f"{cid} {cs[f'{col}4'].value}", cs[f"{col}{rows[cid]}"].value, v,
                  0.0005 if col == "M" else 0.5)
    print("P&L Y1 year total")
    labels = {pl[f"A{r}"].value: r for r in range(6, 16)}
    for label, v in [("Contract revenue", 970000), ("Less raw materials used", 372000),
                     ("Less production wages", 669600), ("Contribution margin", -71600),
                     ("Less fixed operating expenses", 200000), ("Operating profit", -271600)]:
        check(label, pl[f"B{labels[label]}"].value, v, 0.5)
    check("CM %", pl[f"B{labels['CM %']}"].value, -0.0738)
    check("Operating margin %", pl[f"B{labels['Operating margin %']}"].value, -0.2800)
    print("P&L Y1 contract columns")
    for col, cid in zip("CDE", ["C05", "C08", "C09"]):
        check(f"{cid} contribution margin", pl[f"{col}{labels['Contribution margin']}"].value,
              want[cid]["L"], 0.5)
    print("Summary")
    check("Cumulative operating profit after Y1", sm["B11"].value, -271600, 0.5)

    # --- Sensitivity: C09 quantity 30 --------------------------------------
    alt = tmp / "alt.xlsx"
    wb_f = load_workbook(SRC)
    wb_f["Contracts Y1"][f"C{rows['C09']}"] = 30
    wb_f.save(alt)
    wb2 = recalc(alt)
    scan_errors(wb2, "C09 qty 30")
    c2 = wb2["Contracts Y1"]
    print("C09 with accepted qty 30")
    check("C09 production qty", c2[f"E{rows['C09']}"].value, 36, 0)
    check("C09 RM price (Copper 31-40 tier)", c2[f"G{rows['C09']}"].value, 3000, 0)
    # (the delivered workbook was never changed, so C09 is still 40 there)
    check("Delivered C09 qty still 40", load_workbook(SRC)["Contracts Y1"][f"C{rows['C09']}"].value,
          40, 0)

print("\nALL CHECKS PASSED" if not failures else f"\n{len(failures)} FAILURE(S): {failures}")
sys.exit(1 if failures else 0)
